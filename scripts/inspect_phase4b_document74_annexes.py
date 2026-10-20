"""Read official XLSX values/formulas without editing or recalculating source workbooks."""
from pathlib import Path
import hashlib
import json
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
E = ROOT / 'storage/app/phase4b-document74-annex-research-20261014'
source = json.loads((E / 'primary-annex-observation.json').read_text(encoding='utf-8'))
records = []
tokens = ['comExt', 'nFatura', 'vFatura', 'nContCambio', 'vAjusteBCIBSCBSComExt',
          'mdPrestacao', 'nDUIMP', 'nDUE', 'serie', 'indDest', 'dest', 'tomador', '15/10/2026']
for item in source['sources']:
    path = E / item['file']
    assert hashlib.sha256(path.read_bytes()).hexdigest() == item['sha256']
    workbook = load_workbook(path, read_only=True, data_only=False)
    all_rows, selected, sheets = [], [], []
    for sheet in workbook:
        count = 0
        for row in sheet.iter_rows():
            cells = {cell.coordinate: cell.value for cell in row if cell.value is not None}
            if not cells:
                continue
            count += 1
            entry = {'sheet': sheet.title, 'row': row[0].row, 'cells': cells}
            all_rows.append(entry)
            text = ' '.join(str(value) for value in cells.values())
            if any(token in text for token in tokens):
                selected.append(entry)
        sheets.append({'name': sheet.title, 'nonempty_rows': count,
                       'max_row': sheet.max_row, 'max_column': sheet.max_column})
    workbook.close()
    target = E / (path.stem + '-rows.json')
    with target.open('x', encoding='utf-8') as stream:
        json.dump(all_rows, stream, indent=2, ensure_ascii=False, default=str)
    record = {'file': item['file'], 'sha256': item['sha256'], 'sheets': sheets,
              'extracted_rows_file': target.name,
              'extracted_rows_sha256': hashlib.sha256(target.read_bytes()).hexdigest(),
              'selected_rows': selected}
    records.append(record)
    print(item['file'], sheets, 'selected_rows=', len(selected))
nt = ROOT / 'storage/app/phase4b-document74-primary-research-20261014/nt009-v1-01.pdf'
with (E / 'nt009-full-text.txt').open('x', encoding='utf-8') as stream:
    for number, page in enumerate(PdfReader(nt).pages, 1):
        stream.write(f'PHYSICAL PAGE {number}\n{page.extract_text()}\n\n')
with (E / 'annex-content-inspection.json').open('x', encoding='utf-8') as stream:
    json.dump({'status': 'OFFICIAL_ANNEX_VALUES_AND_FORMULAS_EXTRACTED_NO_ACTIVATION',
               'workbooks': records, 'production_ready': False}, stream, indent=2, ensure_ascii=False, default=str)
