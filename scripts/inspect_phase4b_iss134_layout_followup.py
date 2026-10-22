"""Read-only technical workbook follow-up while134 regression inputs remain frozen."""
import hashlib
import json
import os
import re
import urllib.request
from datetime import datetime, timezone
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = Path('D:/ProjetoSaasExports/phase4b-iss134-20261022-v1/layout-followup-v1')
OUT.mkdir(exist_ok=False)


def save(path, raw):
    with path.open('xb') as stream:
        stream.write(raw)
        stream.flush()
        os.fsync(stream.fileno())
    assert hashlib.sha256(path.read_bytes()).digest() == hashlib.sha256(raw).digest()


sources = []
for filename, url in [
    ('rtc-portal.html', 'https://www.gov.br/nfse/pt-br/biblioteca/documentacao-tecnica/rtc'),
    ('nt009-v1-01.pdf', 'https://www.gov.br/nfse/pt-br/biblioteca/documentacao-tecnica/rtc/nota-tecnica-009-se-cgnfs-e-v-1-01.pdf'),
    ('anexovi-v1-04-01.xlsx', 'https://www.gov.br/nfse/pt-br/biblioteca/documentacao-tecnica/rtc/anexovi-leiautesrn_rtc_ibscbs-v1-04-01-nt009.xlsx')]:
    with urllib.request.urlopen(urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'}), timeout=30) as response:
        raw = response.read()
        record = {'file': filename, 'url': url, 'final_url': response.url, 'content_type': response.headers.get('Content-Type'),
                  'observed_at': datetime.now(timezone.utc).isoformat(), 'bytes': len(raw), 'sha256': hashlib.sha256(raw).hexdigest()}
    save(OUT / filename, raw)
    assert raw.startswith(b'PK') if filename.endswith('.xlsx') else raw.startswith(b'%PDF') if filename.endswith('.pdf') else b'<html' in raw.lower()
    sources.append(record)
save(OUT / 'capture.json', json.dumps({'records': sources}, indent=2).encode())
pattern = re.compile(r'cIndOp|410027|cTribMun|exporta', re.I)
analysis = []
for label, path in [('new10401', OUT / 'anexovi-v1-04-01.xlsx'),
                    ('preserved10400', ROOT / 'resources/fiscal/sources/official/anexovi-leiautesrn_rtc_ibscbs-v1-04-00-2013-nt009.xlsx')]:
    workbook = openpyxl.load_workbook(path, data_only=False, keep_links=False)
    matched = []
    sheets = []
    for sheet in workbook.worksheets:
        sheets.append({'sheet': sheet.title, 'rows': sheet.max_row, 'columns': sheet.max_column})
        for row in sheet.iter_rows():
            if not any(isinstance(cell.value, str) and pattern.search(cell.value) for cell in row):
                continue
            cells = []
            for cell in row:
                if cell.value is not None:
                    color = cell.fill.fgColor
                    cells.append({'coordinate': cell.coordinate, 'value': cell.value,
                                  'strike': bool(cell.font.strike), 'fill_type': color.type,
                                  'fill_rgb': color.rgb if color.type == 'rgb' else None})
            matched.append({'sheet': sheet.title, 'row': row[0].row, 'cells': cells})
    analysis.append({'label': label, 'source_sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                     'sheets': sheets, 'matching_rows': matched})
    workbook.close()
save(OUT / 'workbook-review.json', json.dumps({'analysis': analysis, 'read_only': True, 'fiscal_policy_installed': False},
                                            ensure_ascii=False, indent=2, default=str).encode())
summary = {'captured_sources': sources, 'matching_rows': {a['label']: len(a['matching_rows']) for a in analysis},
           'read_only': True, 'fiscal_policy_installed': False}
save(OUT / 'validation.json', json.dumps(summary, indent=2).encode())
print(json.dumps({k: v for k, v in summary.items() if k != 'captured_sources'}))
