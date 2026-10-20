"""Independent source/row checks and accepted101 integrity; research only,no installed gate."""
from pathlib import Path
from datetime import datetime, timezone
from zipfile import ZipFile
import hashlib
import json
import xml.etree.ElementTree as ET
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
E=Path(r'D:\ProjetoSaasExports\phase4b-identity102-primary-preparation-20261016')
A=ROOT/'storage/app/phase4b-catalog101-development-20261016/catalog101-development-validation.json'
F=Path(r'D:\ProjetoSaasExports\phase4b-catalog101-full-suite-20261016')
def read(p): return json.loads(p.read_text(encoding='utf-8'))
def sha(p):
    with p.open('rb') as stream: return hashlib.file_digest(stream,'sha256').hexdigest()
def save(p,v):
    with p.open('x',encoding='utf-8') as stream: json.dump(v,stream,ensure_ascii=False,indent=2)
assert sha(A)=='900bd64cfa7a0eb7775974f936a9d12369929b12028c59fe50e3932a2ed11360'
accepted=read(A)
assert accepted['full_suite']['tests']==6296 and accepted['full_suite']['assertions']==94755
for p,h in accepted['runtime_hashes'].items(): assert sha(ROOT/p)==h,p
frozen=read(F/'preflight.json')['input_hashes']
for p,h in frozen.items(): assert sha(Path(p))==h,p
captures=read(E/'primary-source-observation.json')['sources']+read(E/'annex-observation.json')
for source in captures:
    assert sha(E/source['file'])==source['sha256']
    assert source['tls_validation']=='WINDOWS_DEFAULT_ENABLED'
assert len(captures)==6
assert sha(E/'manual-api.pdf')=='ac2f36e34ff565cc36d67c5d67415c33cc09f27b2dea006e8da9bcd2ddce5581'
assert sha(E/'restricted-xsd.zip')=='6c7e0510d3ecff4454f291f4e10b742d27a4818f23aab181494f96d0ea79f3dc'
inspection=read(E/'source-inspection.json')
ns={'x':'http://www.w3.org/2001/XMLSchema'}
with ZipFile(E/'restricted-xsd.zip') as z:
    for row in inspection['types']:
        assert hashlib.sha256(z.read(row['member'])).hexdigest()==row['member_sha256']
    simple=ET.fromstring(z.read('tiposSimples_v1.01.xsd'))
    pattern=lambda name:simple.find("x:simpleType[@name='"+name+"']/x:restriction/x:pattern",ns).get('value')
    assert pattern('TSIdDPS')=='DPS[0-9]{7}(1[0-9]{14}|2[0-9A-Z]{14})[0-9]{20}'
    assert pattern('TSSerieDPS')=='[0-9]{1,4}|[0-8][0-9]{4}'
    assert pattern('TSNumDPS')=='[1-9]{1}[0-9]{0,14}'
selected=[]
for kind in ['restricted','production']:
    artifact=read(E/f'annex-{kind}-values.json')
    wb=openpyxl.load_workbook(E/f'annex-{kind}.xlsx',read_only=False,data_only=True)
    for name,start,end in [('LEIAUTE DPS_NFS-e ',102,112),('RN DPS_NFS-e',139,156)]:
        sheet=wb[name]; values=next(s['values'] for s in artifact if s['name']==name)
        for row in range(start,end+1):
            native=[sheet.cell(row,col).value for col in range(1, min(sheet.max_column,15)+1)]
            assert native==values[row-1][:len(native)],(kind,name,row)
            selected.append({'environment_source':kind,'sheet':name,'range':f'A{row}:O{row}',
                'values':native,'hidden':bool(sheet.row_dimensions[row].hidden)})
    wb.close()
for kind in ['restricted','production']:
    layout=next(r['values'] for r in selected if r['environment_source']==kind and r['sheet']=='LEIAUTE DPS_NFS-e ' and r['range']=='A106:O106')
    assert '00001 a 49999' in layout[8] and '50000 a 69999' in layout[8]
    rules=[r for r in selected if r['environment_source']==kind and r['sheet']=='RN DPS_NFS-e']
    assert any('E0004' in r['values'] for r in rules)
    assert any('E0010' in r['values'] for r in rules)
save(E/'identity-rows-crosscheck-v2.json',{'selected_rows':selected,'method':'ArtifactTool values vs openpyxl cached values;no workbook edits;column A uses ROW formulas for source line numbering'})
report={'status':'IDENTITY102_PRIMARY_SOURCE_RESEARCH_VERIFIED_NOT_IMPLEMENTED',
    'created_at':datetime.now(timezone.utc).isoformat(),'sources':captures,'primary_captures':6,
    'manual_pages_read':6,'manual_identity_page_visually_reviewed':4,
    'selected_rows_crosschecked':len(selected),'accepted101_inputs_unchanged':len(frozen),
    'accepted101_runtime_hashes_unchanged':len(accepted['runtime_hashes']),
    'acceptance101_sha256':sha(A),'source_inspection_sha256':sha(E/'source-inspection.json'),
    'crosscheck_sha256':sha(E/'identity-rows-crosscheck-v2.json'),
    'auditor_sha256':sha(Path(__file__)),
    'new_contract_required':True,'runtime_integrated':False,'fiscal_identifiers_qualified':False,
    'government_evidence_authenticated':False,'documentary_currency_conversion_qualified':False,
    'per_document_tax_amount_calculated':False,'xml_generated':False,'production_ready':False,
    'next_required':['Source-bound identity component with separate operation/document/material mapping;do not rename historic proofs.',
        'Explicit emitter/channel/environment,series and number boundaries,ID recomposition,distinct sibling identities.',
        'Municipal/CNC/CNPJ authorization and prior-number duplication are separate unavailable remote evidence,not inferred from string validity.',
        'Native cases/tamper/as-of boundaries before versioned integration;preserve101 and all historical replays.']}
save(E/'identity102-research-validation-v2.json',report)
print(json.dumps({k:report[k] for k in ['status','primary_captures','selected_rows_crosschecked','accepted101_inputs_unchanged','accepted101_runtime_hashes_unchanged','runtime_integrated']}),flush=True)
